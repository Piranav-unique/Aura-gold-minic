"""Generate CA ownership certificate Word document for AURUM GOLD AND SILVERS."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


def set_run(run, *, bold=False, size=11):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")


def add_center(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run(run, bold=bold, size=size)
    return p


def add_left(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run(run, bold=bold, size=size)
    return p


def add_mixed(doc, parts, *, size=11, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold in parts:
        run = p.add_run(text)
        set_run(run, bold=bold, size=size)
    return p


def shade_cell(cell, fill="F5E6C8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def main() -> Path:
    out = Path(__file__).resolve().parent / "CA_Certificate_AURUM_GOLD_AND_SILVERS.docx"
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    add_center(doc, "TO WHOMSOEVER IT MAY CONCERN", bold=True, size=14, space_after=4)
    add_center(
        doc,
        "CERTIFICATE OF OWNERSHIP / PROPRIETORSHIP",
        bold=True,
        size=13,
        space_after=12,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta.paragraph_format.space_after = Pt(10)
    r1 = meta.add_run("Date: _______________")
    set_run(r1, size=11)
    meta.add_run("\t\t\t")
    r2 = meta.add_run("Place: Madurai")
    set_run(r2, size=11)

    add_mixed(
        doc,
        [
            ("This is to certify that I/We, ", False),
            ("[CA Full Name]", True),
            (", Chartered Accountant(s), holding Membership No. ", False),
            ("[XXXXXX]", True),
            (", in practice under the firm name ", False),
            ("[CA Firm Name]", True),
            (", Firm Registration No. (FRN) ", False),
            ("[XXXXXX]", True),
            (
                ", have examined the books of accounts, GST / tax records, "
                "and other relevant documents of the following business concern:",
                False,
            ),
        ],
        space_after=12,
    )

    add_left(doc, "BUSINESS DETAILS", bold=True, size=12, space_after=8)

    table = doc.add_table(rows=12, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    rows = [
        ("Particulars", "Details"),
        ("Name of the Business / Trade Name", "AURUM GOLD AND SILVERS"),
        ("Also known as / Brand name", "AGS Gold / AGS Digi Gold"),
        ("Nature of Constitution", "Sole Proprietorship"),
        (
            "Nature of Business",
            "Jewellery retail – sale of gold and silver jewellery, coins, "
            "and ornaments; and operation of a gold savings / advance booking "
            "scheme through the mobile application AGS Digi Gold "
            "(Android package: com.agsgold.ags_gold)",
        ),
        (
            "Principal Place of Business",
            "45 A, South Pandiyan Agil Street, Madurai, Tamil Nadu – 625001",
        ),
        ("Proprietor’s Full Name", "VIJAYAKUMAR MANICKAM"),
        ("Proprietor’s PAN", "AUGPV6458B"),
        (
            "Business / Firm PAN",
            "AUGPV6458B (same as proprietor for sole proprietorship)",
        ),
        ("GSTIN", "[Your GSTIN]"),
        ("Shop / Trade License No.", "[If available]"),
        (
            "Registered Email / Mobile / Website",
            "aurumgoldsilvers@gmail.com | +91 99437 95005 | https://aurumgold.co.in",
        ),
    ]

    for i, (left, right) in enumerate(rows):
        cell0, cell1 = table.rows[i].cells
        cell0.text = left
        cell1.text = right
        for cell in (cell0, cell1):
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_run(run, bold=(i == 0), size=10)
        if i == 0:
            shade_cell(cell0)
            shade_cell(cell1)
        if i in (1, 6):
            for run in cell1.paragraphs[0].runs:
                set_run(run, bold=True, size=10)
        cell0.width = Inches(2.6)
        cell1.width = Inches(4.2)

    doc.add_paragraph()
    add_left(doc, "CERTIFICATION", bold=True, size=12, space_after=8)
    add_left(
        doc,
        "Based on verification of the documents and information produced "
        "before us, we hereby certify that:",
        space_after=8,
    )

    points = [
        "AURUM GOLD AND SILVERS is a Sole Proprietorship concern.",
        "The sole proprietor of the said concern is VIJAYAKUMAR MANICKAM, "
        "holding PAN AUGPV6458B.",
        "The ownership of the business rests entirely with the above-named "
        "proprietor, and no other person has any ownership rights in the firm.",
        "The business is engaged in physical jewellery retail (gold and silver "
        "jewellery / coins / ornaments) and related customer savings / advance "
        "booking through its own application. It is not a third-party "
        "digital-gold vaulting platform (such as SafeGold / MMTC-PAMP type vaulting).",
        "The mobile application AGS Digi Gold is owned and operated by "
        "AURUM GOLD AND SILVERS for its own jewellery retail / gold savings "
        "scheme business.",
        "This certificate is issued at the request of the proprietor for the "
        "purpose of: [Google Play / Razorpay / Bank / NSWS / Payment Gateway / "
        "Other – please specify] verification and related compliance.",
    ]

    for i, text in enumerate(points, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(f"{i}. {text}")
        set_run(run, size=11)

    doc.add_paragraph()
    add_left(doc, "DISCLAIMER", bold=True, size=12, space_after=6)
    add_left(
        doc,
        "This certificate is issued on the basis of documents and information "
        "provided by the proprietor and examination of relevant records available "
        "up to the date of this certificate. We have not conducted a statutory "
        "audit for the purpose of this certificate unless separately engaged. "
        "This certificate is not a guarantee of future performance or solvency.",
        space_after=14,
    )

    add_left(doc, "FOR / ON BEHALF OF", bold=True, size=11, space_after=2)
    add_left(doc, "[CA Firm Name]", bold=True, size=11, space_after=2)
    add_left(doc, "Chartered Accountants", size=11, space_after=18)
    add_left(doc, "______________________________", size=11, space_after=2)
    add_left(
        doc,
        "Signature & Seal of Chartered Accountant",
        bold=True,
        size=11,
        space_after=12,
    )

    sign_table = doc.add_table(rows=7, cols=2)
    sign_table.style = "Table Grid"
    sign_rows = [
        ("Name of CA", "[CA Full Name]"),
        ("Membership No.", "[XXXXXX]"),
        ("Firm Name", "[CA Firm Name]"),
        ("FRN", "[XXXXXX]"),
        ("UDIN", "[________________]"),
        ("Date", "[DD/MM/YYYY]"),
        ("Place", "Madurai"),
    ]
    for i, (left, right) in enumerate(sign_rows):
        sign_table.rows[i].cells[0].text = left
        sign_table.rows[i].cells[1].text = right
        for run in sign_table.rows[i].cells[0].paragraphs[0].runs:
            set_run(run, bold=True, size=10)
        for run in sign_table.rows[i].cells[1].paragraphs[0].runs:
            set_run(run, size=10)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    run = note.add_run("Note for Proprietor / CA: ")
    set_run(run, bold=True, size=10)
    run2 = note.add_run(
        "Please replace bracketed fields such as GSTIN, Shop License, CA details, "
        "and purpose. Use the exact trade name as per GST certificate "
        "(AURUM GOLD AND SILVERS or AURUM GOLD & SILVERS). CA must print on "
        "letterhead, sign, stamp, and generate UDIN on the ICAI portal before submission."
    )
    set_run(run2, size=10)

    doc.save(out)
    return out


if __name__ == "__main__":
    path = main()
    print(path)
